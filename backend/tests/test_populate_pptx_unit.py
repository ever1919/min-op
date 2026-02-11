"""Unit tests for PPTX rendering (populate_pptx)."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation

import populate_pptx


def _find_shape_by_name(slide, name: str):
    """Return the first shape whose name matches (case-insensitive)."""
    target = name.strip().lower()
    for shape in slide.shapes:
        if getattr(shape, "name", "").strip().lower() == target:
            return shape
    return None


def test_generates_file(template_digicore: Path, backend_cwd: Path):
    """populate_pptx() creates a non-empty PPTX file."""
    df = pd.DataFrame(
        [
            {"section_name": "NAME", "output": "John Doe"},
            {"section_name": "TOWER", "output": "Control Tower"},
            {"section_name": "PROFILE OVERVIEW", "output": "Experienced analyst."},
            {"section_name": "ROLE 1", "output": "Consulting Analyst\nDid X\nDid Y"},
        ]
    )

    out_path = populate_pptx.populate_pptx(df, coe_selected="Digi Core")
    out_file = backend_cwd / out_path

    assert out_file.exists(), f"Expected PPTX output to exist: {out_file}"
    assert out_file.stat().st_size > 0


def test_name_uppercase(template_digicore: Path, backend_cwd: Path):
    """NAME section is rendered in uppercase."""
    df = pd.DataFrame([{"section_name": "NAME", "output": "John Doe"}])

    out_path = populate_pptx.populate_pptx(df, coe_selected="Digi Core")
    prs = Presentation(backend_cwd / out_path)
    slide = prs.slides[0]

    name_shape = _find_shape_by_name(slide, "NAME")
    assert name_shape is not None

    assert name_shape.text_frame.text.strip() == "JOHN DOE"
